"""Local, bounded file previews. No model calls or external document viewers."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile

TEXT_LIMIT = 120_000


def preview_info(path: Path, parsed_text=None):
    suffix = path.suffix.lower()
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}:
        from PIL import Image

        with Image.open(path) as img:
            return {"kind": "image", "width": img.width, "height": img.height}
    if suffix == ".pdf":
        import fitz

        with fitz.open(path) as document:
            if document.needs_pass:
                return {
                    "kind": "unsupported",
                    "note": "文件已加密，请下载后输入密码查看。",
                }
            return {"kind": "pdf", "page_count": len(document)}
    if suffix in {".txt", ".md"}:
        with path.open("rb") as stream:
            raw = stream.read(TEXT_LIMIT * 4 + 1)
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                import codecs

                text = codecs.getincrementaldecoder(encoding)().decode(
                    raw, final=len(raw) <= TEXT_LIMIT * 4
                )
                break
            except UnicodeDecodeError:
                continue
        else:
            text = raw.decode("utf-8", errors="replace")
        return text_preview(text, truncated=len(raw) > TEXT_LIMIT * 4)
    if suffix in {".docx", ".pptx"}:
        # Avoid unbounded decompression in a synchronous read-only request.
        if path.stat().st_size > 64 * 1024 * 1024:
            return {
                "kind": "unsupported",
                "note": "文件较大，暂不提供文字预览，请下载原文件查看。",
            }
        try:
            with ZipFile(path) as archive:
                if sum(i.file_size for i in archive.infolist()) > 128 * 1024 * 1024:
                    return {
                        "kind": "unsupported",
                        "note": "文件内容较大，请下载原文件查看。",
                    }
        except BadZipFile:
            return {
                "kind": "unsupported",
                "note": "文件格式异常，无法预览。可以下载检查原文件。",
            }
        chunks, length = [], 0

        def add(value):
            nonlocal length
            chunks.append(value)
            length += len(value)

        if suffix == ".docx":
            from xml.etree import ElementTree

            with ZipFile(path) as archive:
                if archive.getinfo("word/document.xml").file_size > 8 * 1024 * 1024:
                    return {
                        "kind": "unsupported",
                        "note": "正文内容较大，请下载原文件查看。",
                    }
                root = ElementTree.fromstring(archive.read("word/document.xml"))
            namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            for paragraph in root.iter(namespace + "p"):
                add(
                    "".join(node.text or "" for node in paragraph.iter(namespace + "t"))
                )
                if length > TEXT_LIMIT:
                    break
        else:
            from pptx import Presentation

            presentation = Presentation(path)
            for i, slide in enumerate(presentation.slides, 1):
                add(f"第 {i} 页")
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        add(shape.text)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            add(" | ".join(cell.text for cell in row.cells))
                    if length > TEXT_LIMIT:
                        break
                if length > TEXT_LIMIT:
                    break
        return text_preview(
            "\n\n".join(chunks),
            note="文字内容预览，不还原原始排版；图片中的文字不会自动识别。",
        )
    if parsed_text:
        return text_preview(parsed_text, note="已解析的文字内容预览，不还原原始排版。")
    return {"kind": "unsupported", "note": "此格式暂不支持在线预览，可下载原文件查看。"}


def text_preview(text, note="", truncated=False):
    return {
        "kind": "text",
        "text": text[:TEXT_LIMIT],
        "truncated": truncated or len(text) > TEXT_LIMIT,
        "note": note or ("文件没有可预览的文字。" if not text.strip() else ""),
    }


def pdf_page(path: Path, page: int):
    import fitz

    with fitz.open(path) as document:
        if document.needs_pass or not 1 <= page <= len(document):
            raise ValueError("页码无效或文件已加密")
        source = document[page - 1]
        longest = max(source.rect.width, source.rect.height)
        if longest <= 0:
            raise ValueError("页面尺寸无效")
        pixmap = source.get_pixmap(
            matrix=fitz.Matrix(min(1.6, 1800 / longest), min(1.6, 1800 / longest)),
            alpha=False,
        )
        return BytesIO(pixmap.tobytes("png"))
