"""Strict semantic adapter over the existing python-pptx/PPTXBuilder stack."""
from io import BytesIO
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Pt
from pptx.opc.package import Part
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.xmlchemy import OxmlElement
from utils.pptx_builder import PPTXBuilder
from .model import load_page

EMU = 9525  # CSS px at 96 dpi
SVG_NS = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'


def coords(box):
    return tuple(round(v*EMU) for v in (box.x, box.y, box.w, box.h))


def set_font(run, style):
    run.font.name = style.font
    run.font.size = Pt(style.size*.75)
    run.font.bold, run.font.italic = style.bold, style.italic
    run.font.color.rgb = RGBColor.from_string(style.color)
    # Explicit east Asian typeface avoids theme font substitution for Chinese.
    rpr = run._r.get_or_add_rPr()
    ea = OxmlElement('a:ea')
    ea.set('typeface', style.font)
    rpr.append(ea)


def add_svg(slide, picture, raw):
    package = slide.part.package
    part = Part(package.next_partname('/ppt/media/semantic%d.svg'), 'image/svg+xml', package, raw)
    rid = slide.part.relate_to(part, RT.IMAGE)
    blip = picture._element.blipFill.blip
    extlist = OxmlElement('a:extLst')
    ext = OxmlElement('a:ext')
    ext.set('uri', '{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    from lxml import etree
    svg = etree.SubElement(ext, '{'+SVG_NS+'}svgBlip', nsmap={'asvg': SVG_NS})
    svg.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed', rid)
    extlist.append(ext)
    blip.append(extlist)


def build_deck(pages, store):
    if not pages:
        raise ValueError('empty semantic deck')
    pages = [load_page(p.model_dump(mode='json')) for p in pages]
    if len({p.id for p in pages}) != len(pages):
        raise ValueError('duplicate page ID')
    if len({(p.width, p.height) for p in pages}) != 1:
        raise ValueError('PPTX requires one slide size per deck')
    first = pages[0]
    builder = PPTXBuilder(first.width/96, first.height/96)
    builder.create_presentation()
    for page in pages:
        if page.status != 'ready':
            raise ValueError(f'{page.id}: needs correction')
        store.verify(page)
        slide = builder.add_blank_slide()
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(page.background)
        assets = {a.id: a for a in page.assets}
        for node in sorted(page.nodes, key=lambda n: n.layer):
            xywh = coords(node.box)
            if node.kind == 'text':
                sh = slide.shapes.add_textbox(*xywh)
                tf = sh.text_frame
                tf.clear()
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                tf.word_wrap = True
                tf.auto_size = MSO_AUTO_SIZE.NONE
                for i, para in enumerate(node.text.paragraphs):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[para.align]
                    p.line_spacing = para.line_spacing
                    p.space_before = Pt(0)
                    p.space_after = Pt(para.space_after*.75)
                    for style in para.runs:
                        run = p.add_run()
                        run.text = style.text
                        set_font(run, style)
            elif node.kind == 'shape':
                spec = node.shape
                sh = slide.shapes.add_shape({'rect': MSO_SHAPE.RECTANGLE, 'roundRect': MSO_SHAPE.ROUNDED_RECTANGLE, 'ellipse': MSO_SHAPE.OVAL, 'rightArrow': MSO_SHAPE.RIGHT_ARROW}[spec.geometry], *xywh)
                for effect in sh._element.xpath('./p:style/a:effectRef'):
                    effect.set('idx', '0')
                if spec.geometry == 'roundRect':
                    sh.adjustments[0] = 0.04
                sh.fill.solid()
                sh.fill.fore_color.rgb = RGBColor.from_string(spec.fill)
                if spec.stroke:
                    sh.line.color.rgb = RGBColor.from_string(spec.stroke)
                    sh.line.width = Pt(spec.stroke_width*.75)
                else:
                    sh.line.fill.background()
            elif node.kind in ('image', 'svg'):
                pic = node.picture
                asset = assets[pic.fallback_asset_id if node.kind == 'svg' else pic.asset_id]
                sh = slide.shapes.add_picture(BytesIO(store.read(asset)), *xywh)
                sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom = (pic.crop.left, pic.crop.top, pic.crop.right, pic.crop.bottom)
                if node.kind == 'svg':
                    add_svg(slide, sh, store.read(assets[pic.asset_id]))
            else:
                spec = node.table
                sh = slide.shapes.add_table(len(spec.rows), len(spec.column_weights), *xywh)
                table = sh.table
                widths = [round(xywh[2]*w/sum(spec.column_weights)) for w in spec.column_weights]
                widths[-1] += xywh[2]-sum(widths)
                for col, width in zip(table.columns, widths):
                    col.width = width
                from .model import Run
                for r, row in enumerate(spec.rows):
                    for c, value in enumerate(row):
                        cell = table.cell(r, c)
                        cell.margin_left = cell.margin_right = round(10*EMU)
                        cell.margin_top = cell.margin_bottom = round(6*EMU)
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(spec.header_fill if r == 0 else spec.body_fill)
                        cell.text_frame.clear()
                        run = cell.text_frame.paragraphs[0].add_run()
                        run.text = value
                        set_font(run, Run(text=value, font=spec.font, size=spec.size, color=spec.header_color if r == 0 else spec.text_color, bold=r == 0))
            sh.name = f'{node.id}::{node.name}'
            sh.rotation = node.rotation
            sh._element.xpath('.//p:cNvPr')[0].set('descr', node.name)
        # Groups are built bottom-up, only from contiguous ordered siblings.
        # Preserve page-absolute child geometry using off == chOff, ext == chExt.
        elements = {sh.name.split('::', 1)[0]: sh for sh in slide.shapes}
        def depth(g):
            return len(page.ancestors(g))
        for group in sorted(page.groups, key=depth, reverse=True):
            members = [o for o in page.nodes+page.groups if o.parent_id == group.id]
            members.sort(key=lambda o: min(n.layer for n in page.nodes if n.id == o.id or o.id in page.ancestors(n)))
            shapes = [elements[m.id] for m in members]
            tree = slide.shapes._spTree
            index = list(tree).index(shapes[0]._element)
            gs = slide.shapes.add_group_shape(shapes)
            gs.name = f'{group.id}::{group.name}'
            gs._element.xpath('./p:nvGrpSpPr/p:cNvPr')[0].set('descr', group.name)
            tree.remove(gs._element)
            tree.insert(index, gs._element)
            elements[group.id] = gs
        slide.notes_slide.notes_text_frame.text = '\n'.join([
            f'Semantic schema {page.schema_version}; page {page.id}@{page.revision}; source {page.source.id}@{page.source.revision}',
            *page.issues(),
        ])
    stream = BytesIO()
    builder.prs.save(stream)
    return stream.getvalue()
