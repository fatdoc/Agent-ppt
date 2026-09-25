import base64
import copy
import hashlib
import json
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from PIL import Image
from pptx import Presentation
from services.semantic_export.model import load_page, digest, correct_node, Node
from services.semantic_export import AssetStore, export_pages, export_jobs, PageJob, SemanticExportFailure
from services.semantic_export.builder import build_deck
from services.semantic_export.audit import audit_pptx
from services.semantic_export.adapters import merge_paragraphs, from_editable_image
from services.image_editability.data_models import EditableImage, EditableElement, BBox
from scripts.semantic_export.sample import samples, basic_page, node, text


@pytest.fixture(autouse=True)
def offline_only(monkeypatch):
    import socket
    def blocked(*args, **kwargs):
        raise AssertionError('Network access forbidden in semantic export tests')
    monkeypatch.setattr(socket.socket, 'connect', blocked)


@pytest.fixture
def fixture(tmp_path):
    pages = samples(tmp_path)
    return pages, AssetStore(tmp_path,user_id='sample-user',project_id='sample-project')


def test_serialization_version_and_unknown_fields(fixture):
    page=fixture[0][0]
    assert load_page(json.loads(page.model_dump_json())) == page
    for change in ({'schema_version':'2.0'},{'schema_version':None},{'browser_json':{}},{'unit':'inches'}):
        with pytest.raises(ValueError):
            load_page({**page.model_dump(mode='json'),**change})


@pytest.mark.parametrize('mutation',[
    lambda d:d['nodes'].append(d['nodes'][0]),
    lambda d:d['nodes'][0]['box'].update(w=-1),
    lambda d:d['nodes'][0]['box'].update(x=float('nan')),
    lambda d:d['nodes'][0].update(rotation=90),
    lambda d:d['nodes'][0].update(parent_id='missing'),
    lambda d:d['groups'][0].update(parent_id='content-group'),
    lambda d:d['nodes'][2].update(layer=20),
    lambda d:d['assets'][0].update(project_id='other-project'),
    lambda d:d['assets'][0].update(uri='../other/file.png'),
    lambda d:d['assets'][0].update(uri='/tmp/photo.png'),
    lambda d:d['assets'][0].update(uri='a%2fb.png'),
    lambda d:d['nodes'][4]['picture'].update(asset_id='missing'),
    lambda d:d['nodes'][2]['picture'].update(fallback_asset_id=None),
    lambda d:d['nodes'][4]['picture'].update(preserve_whole=False),
    lambda d:d['nodes'][4]['picture'].update(crop={'left':.1}),
    lambda d:d['nodes'][4]['box'].update(x=0,y=0,w=1280,h=720),
    lambda d:d['nodes'][3]['box'].update(x=720,y=200,w=200,h=100),
    lambda d:d['assets'][0].update(sha256=d['source']['sha256']),
])
def test_invalid_contracts_are_rejected(fixture,mutation):
    data=fixture[0][0].model_dump(mode='json')
    mutation(data)
    with pytest.raises(ValueError): load_page(data)


def test_low_confidence_and_draft_mismatch_retained(fixture):
    data=fixture[0][0].model_dump(mode='json')
    data['nodes'][3]['provenance'].update(confidence=.2,text_alignment='mismatch',proposed_text='wrong draft')
    page=load_page(data)
    assert len(page.issues())==2
    raw=build_deck([page],fixture[1])
    assert page.nodes[3].text.paragraphs[0].runs[0].text in Presentation(BytesIO(raw)).slides[0].shapes[1].shapes[2].text


def test_whole_native_objects_groups_coordinates_and_relationships(fixture):
    pages,store=fixture
    raw=build_deck(pages,store)
    report=audit_pptx(raw,pages)
    assert report[0]['leaf_objects']==6
    assert report[0]['groups']==1
    assert report[1]['counts']['table']==1
    prs=Presentation(BytesIO(raw))
    assert len(prs.slides[0].shapes[1].shapes)==4
    assert prs.slides[1].shapes[2].has_table
    assert prs.slides[1].shapes[2].table.cell(2,0).text=='卡片模块'


def test_nested_groups_and_crop_rotation(fixture):
    pages,store=fixture
    d=pages[0].model_dump(mode='json')
    d['groups'].append(dict(id='inner',name='nested',module_id='content',parent_id='content-group'))
    for n in d['nodes'][2:4]: n['parent_id']='inner'
    d['nodes'][4]['picture'].update(role='photo',preserve_whole=False,crop=dict(left=.1,right=.15,top=.1,bottom=.05))
    d['nodes'][2]['rotation']=15
    page=load_page(d)
    report=audit_pptx(build_deck([page],store),[page])
    assert report[0]['groups']==2


def test_paragraph_boundaries_and_exact_content():
    a=Node.model_validate(node('a','text',[20,20,200,30],0,'card1',text=text('first')))
    b=Node.model_validate(node('b','text',[20,55,200,30],1,'card1',text=text('second')))
    assert [p.runs[0].text for p in merge_paragraphs([b,a]).text.paragraphs]==['first','second']
    for change in ({'module_id':'card2'},{'column_id':'right'},{'parent_id':'other'},{'rotation':1}):
        changed=Node.model_validate({**b.model_dump(),**change})
        with pytest.raises(ValueError): merge_paragraphs([a,changed])


def test_editable_image_bridge_retains_screenshot_child_text(fixture):
    page=fixture[0][0]
    pic=page.nodes[4]
    b=pic.box
    box=BBox(b.x,b.y,b.x+b.w,b.y+b.h)
    child=EditableElement('internal-text','text',BBox(1,1,20,20),BBox(b.x+1,b.y+1,b.x+20,b.y+20),'tiny')
    image=EditableImage('root','unused-original.png',1280,720,elements=[EditableElement(pic.id,'image',box,box,children=[child])])
    meta=page.model_dump(mode='json')
    meta['nodes']=[n for n in meta['nodes'] if n['id']!=pic.id]
    result=from_editable_image(image,page_metadata=meta,annotations={pic.id:pic.model_dump(mode='json')},assets=[a.model_dump(mode='json') for a in page.assets])
    assert len(result.nodes)==len(page.nodes)
    assert 'internal-text' not in [n.id for n in result.nodes]
    assert any('internal-text' in w for w in result.warnings)
    with pytest.raises(ValueError,match='unassigned'):
        from_editable_image(image,page_metadata=meta,annotations={},assets=meta['assets'])


def test_asset_hash_symlink_tenant_and_active_svg(fixture,tmp_path):
    page,store=fixture[0][0],fixture[1]
    with pytest.raises(ValueError,match='scope'):
        AssetStore(tmp_path,user_id='another',project_id=page.project_id).verify(page)
    asset=page.assets[0]
    path=tmp_path/asset.uri
    path.write_bytes(b'bad')
    with pytest.raises(ValueError,match='hash'): store.read(asset)
    outside=tmp_path.parent/'outside.png'
    outside.write_bytes(b'bad')
    path.unlink();path.symlink_to(outside)
    with pytest.raises(ValueError,match='escapes'): store.read(asset)
    svg=page.assets[1].model_copy(deep=True)
    svgpath=tmp_path/svg.uri
    raw=b'<svg xmlns="http://www.w3.org/2000/svg"><image href="https://invalid.example/x"/></svg>'
    svgpath.write_bytes(raw);svg.sha256=hashlib.sha256(raw).hexdigest()
    with pytest.raises(ValueError,match='external'): store.read(svg)


def test_structural_audit_detects_actual_layer_and_coordinate_corruption(fixture):
    pages,store=fixture
    raw=build_deck(pages,store)
    src=ZipFile(BytesIO(raw))
    for old,new in [(b'name="body::body"',b'name="title::title"'),(b'x="876300"',b'x="876399"')]:
        dest=BytesIO()
        with ZipFile(dest,'w',ZIP_DEFLATED) as z:
            for name in src.namelist():
                content=src.read(name)
                if name=='ppt/slides/slide1.xml':
                    assert old in content
                    content=content.replace(old,new)
                z.writestr(name,content)
        with pytest.raises(ValueError): audit_pptx(dest.getvalue(),pages)


def test_cache_retry_local_correction_and_no_partial_output(fixture,tmp_path):
    pages,store=fixture
    calls=[]
    fail={'business-table'}
    def job(page):
        def extract():
            calls.append(page.id)
            if page.id in fail: raise RuntimeError('fake extractor failed')
            return page.model_dump(mode='json')
        return PageJob(page.id,page.source,'fake-v1','organizer-v1','r1',extract,lambda d:d)
    jobs=[job(p) for p in pages]
    events=[]
    options=dict(store=store,checkpoint_dir=tmp_path/'cache',mode='semantic',on_event=events.append)
    with pytest.raises(SemanticExportFailure) as err: export_jobs(jobs,**options)
    assert err.value.completed==['representative']
    assert err.value.errors['business-table']['stage']=='extract'
    calls.clear();fail.clear();events.clear()
    raw,_=export_jobs(jobs,**options)
    assert calls==['business-table']
    assert all(e['cache_hit'] for e in events if e['page_id']=='representative')
    calls.clear();events.clear()
    replacement=pages[0].nodes[3].model_copy(deep=True)
    replacement.text.paragraphs[0].runs[0].text='局部人工修正，其他页面继续复用。'
    corrected=correct_node(pages[0],replacement.id,replacement,'r2','reviewed text')
    jobs[0]=replace(jobs[0],correction_revision='r2',organize=lambda d:corrected.model_dump(mode='json'))
    export_jobs(jobs,**options)
    assert calls==[]
    assert [e['cache_hit'] for e in events if e['page_id']=='representative']==[True,False,False,False]
    assert all(e['cache_hit'] for e in events if e['page_id']=='business-table')
    assert corrected.corrections[0].before_sha256 != corrected.corrections[0].after_sha256


def test_object_failure_retries_and_corrupt_cache_recovers(fixture,tmp_path,monkeypatch):
    import services.semantic_export.pipeline as pipeline
    pages,store=fixture
    real=pipeline.build_deck
    failed={pages[1].id}
    def builder(pages,store):
        if len(pages)==1 and pages[0].id in failed: raise RuntimeError('fake build failure')
        return real(pages,store)
    monkeypatch.setattr(pipeline,'build_deck',builder)
    options=dict(store=store,checkpoint_dir=tmp_path/'cache',mode='semantic')
    with pytest.raises(SemanticExportFailure) as err: export_pages(pages,**options)
    assert err.value.errors['business-table']['stage']=='objects'
    failed.clear();events=[]
    export_pages(pages,**options,on_event=events.append)
    assert [e['cache_hit'] for e in events if e['page_id']=='business-table']==[True,True,False,False]
    for f in (tmp_path/'cache').rglob('*.json'): f.write_text('broken')
    export_pages(pages,**options)


def test_pending_pages_and_wrong_mode_do_not_silently_succeed(fixture,tmp_path):
    pages,store=fixture
    pages[0].status='needs_correction'
    with pytest.raises(SemanticExportFailure):
        export_pages(pages,store=store,checkpoint_dir=tmp_path/'cache',mode='semantic')
    with pytest.raises(ValueError,match='explicit'):
        export_pages(pages,store=store,checkpoint_dir=tmp_path/'cache',mode='legacy')


def test_table_rejects_ragged_rows(fixture):
    data=fixture[0][1].model_dump(mode='json')
    data['nodes'][2]['table']['rows'][0].pop()
    with pytest.raises(ValueError,match='rectangular'): load_page(data)


def test_cancellation_propagates_without_running_more_pages(fixture,tmp_path):
    from services.semantic_export.pipeline import ExportInterrupted
    pages,store=fixture
    calls=[]
    def cancel():
        calls.append('cancel')
        raise RuntimeError('cancelled')
    with pytest.raises(ExportInterrupted):
        export_pages(pages,store=store,checkpoint_dir=tmp_path/'cache',mode='semantic',check_cancelled=cancel)
    assert calls==['cancel']
    assert not list((tmp_path/'cache').rglob('*.json'))


def test_source_version_and_extractor_invalidation(fixture,tmp_path):
    pages,store=fixture
    page=pages[0]
    job=PageJob(page.id,page.source,'fake-v1','organizer','r1',lambda:page.model_dump(mode='json'),lambda d:d)
    options=dict(store=store,checkpoint_dir=tmp_path/'cache',mode='semantic')
    export_jobs([job],**options)
    events=[]
    export_jobs([replace(job,extraction_signature='fake-v2')],**options,on_event=events.append)
    assert not any(e['cache_hit'] for e in events)
    changed=page.model_copy(deep=True)
    changed.source.revision='source-r2'
    events=[]
    export_jobs([replace(job,source=changed.source,extract=lambda:changed.model_dump(mode='json'))],**options,on_event=events.append)
    assert not any(e['cache_hit'] for e in events)


def test_preserved_screenshot_cannot_be_stretched(fixture):
    page=fixture[0][0].model_copy(deep=True)
    page.nodes[4].box.w*=1.5
    with pytest.raises(ValueError,match='aspect ratio'):
        build_deck([page],fixture[1])


def test_acceptance_failure_retries_without_rebuilding_successes(fixture,tmp_path,monkeypatch):
    import services.semantic_export.pipeline as pipeline
    pages,store=fixture
    real=pipeline.audit_pptx
    failed={pages[1].id}
    def audit(raw,pages):
        if len(pages)==1 and pages[0].id in failed: raise ValueError('fake acceptance failure')
        return real(raw,pages)
    monkeypatch.setattr(pipeline,'audit_pptx',audit)
    options=dict(store=store,checkpoint_dir=tmp_path/'cache',mode='semantic')
    with pytest.raises(SemanticExportFailure) as err: export_pages(pages,**options)
    assert err.value.errors['business-table']['stage']=='acceptance'
    failed.clear();events=[]
    export_pages(pages,**options,on_event=events.append)
    assert [e['cache_hit'] for e in events if e['page_id']=='business-table']==[True,True,True,False]
    assert all(e['cache_hit'] for e in events if e['page_id']=='representative')


def test_legacy_image_export_remains_available(fixture,tmp_path):
    from services.export_service import ExportService
    page,store=fixture[0][0],fixture[1]
    raw=ExportService.create_pptx_from_images([str(tmp_path/page.assets[0].uri)])
    prs=Presentation(BytesIO(raw))
    assert len(prs.slides)==1
    assert len(prs.slides[0].shapes)==1
