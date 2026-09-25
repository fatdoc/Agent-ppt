import copy
import pytest
from models import User,Project,Page,db
from models.editor_document import EditorRevision,EditorDocument
from utils.auth import create_auth_token
from scripts.semantic_export.sample import basic_page,node,text

@pytest.fixture
def case(client,app,monkeypatch):
    monkeypatch.setenv('AUTH_REQUIRED','true')
    with app.app_context():
        owner=User(username='editor-owner');owner.set_password('test-password')
        other=User(username='editor-other');other.set_password('test-password')
        db.session.add_all([owner,other]);db.session.flush()
        project=Project(user_id=owner.id,creation_type='idea');db.session.add(project);db.session.flush()
        page=Page(project_id=project.id,order_index=0);db.session.add(page);db.session.commit()
        raw=basic_page(page.id,[node('title','text',[60,40,900,80],0,text=text('中文语义编辑'))])
        raw.update(user_id=owner.id,project_id=project.id)
        return dict(url=f'/api/projects/{project.id}/editor-document',owner={'Authorization':'Bearer '+create_auth_token(owner)},other={'Authorization':'Bearer '+create_auth_token(other)},raw=raw,project_id=project.id)

def initialize(client,c):
    r=client.post(c['url']+'/initialize',json={'base_revision':0,'pages':[c['raw']]},headers=c['owner'])
    assert r.status_code==200,r.json
    return r.json['data']

def test_save_refresh_conflict_restore_and_immutable_history(client,case):
    c=case;d=initialize(client,c)
    assert d['revision']==1
    d['slides'][0]['elements'][0]['content']=d['slides'][0]['elements'][0]['content'].replace('中文','修改')
    body={'base_revision':1,'slides':d['slides']}
    r=client.put(c['url'],json=body,headers=c['owner']);assert r.status_code==200,r.json
    assert r.json['data']['revision']==2
    assert client.put(c['url'],json=body,headers=c['owner']).status_code==409
    assert client.get(c['url'],headers=c['owner']).json['data']['pages'][0]['nodes'][0]['text']['paragraphs'][0]['runs'][0]['text']=='修改语义编辑'
    r=client.post(c['url']+'/restore',json={'base_revision':2,'revision':1},headers=c['owner']);assert r.status_code==200,r.json
    assert r.json['data']['revision']==3
    assert r.json['data']['pages'][0]['nodes'][0]['text']['paragraphs'][0]['runs'][0]['text']=='中文语义编辑'
    assert client.post(c['url']+'/restore',json={'base_revision':2,'revision':1},headers=c['owner']).status_code==409
    assert EditorRevision.query.filter_by(project_id=c['project_id']).count()==3
    assert EditorRevision.query.filter_by(project_id=c['project_id'],revision=3).one().restored_from_revision==1

def test_scope_auth_size_and_explicit_initialization(client,case):
    c=case
    assert client.put(c['url'],json={'base_revision':0,'pages':[c['raw']]},headers=c['owner']).status_code==404
    initialize(client,c)
    assert client.post(c['url']+'/initialize',json={'base_revision':0,'pages':[c['raw']]},headers=c['owner']).status_code==409
    for suffix in ('','/revisions'):
        assert client.get(c['url']+suffix,headers=c['other']).status_code==404
        assert client.get(c['url']+suffix).status_code==401
    assert client.post(c['url']+'/restore',json={'base_revision':1,'revision':1},headers=c['other']).status_code==404
    assert client.put(c['url'],data='x'*(4*1024*1024+1),headers=c['owner']).status_code==413

def test_invalid_save_does_not_advance_revision(client,case):
    c=case;d=initialize(client,c);e=copy.deepcopy(d['slides'])
    e[0]['elements'][0]['content']='<p><script>bad()</script></p>'
    assert client.put(c['url'],json={'base_revision':1,'slides':e},headers=c['owner']).status_code==400
    assert client.get(c['url'],headers=c['owner']).json['data']['revision']==1
    assert EditorRevision.query.filter_by(project_id=c['project_id']).count()==1
    assert client.put(c['url'],json={'base_revision':1,'slides':[]},headers=c['owner']).status_code==400

def test_editor_never_refreshes_oauth_or_constructs_provider(client,case,monkeypatch,app):
    from models import Settings
    from services.provider_config import active_provider_snapshot
    import services.ai_providers as providers
    def forbidden(*a,**kw):raise AssertionError('editor must not call OAuth or Provider')
    monkeypatch.setattr(Settings,'get_openai_oauth_token',forbidden)
    monkeypatch.setattr(providers,'get_text_provider',forbidden)
    monkeypatch.setattr(providers,'get_image_provider',forbidden)
    monkeypatch.setenv('GOOGLE_API_KEY','fake-default-to-detect-fallback')
    c=case
    with app.app_context():
        user=Project.query.filter_by(id=c['project_id']).one().user_id
        db.session.add(Settings(user_id=user,ai_provider_format='codex'));db.session.commit()
    d=initialize(client,c)
    assert client.get(c['url'],headers=c['owner']).status_code==200
    assert client.put(c['url'],json={'base_revision':1,'slides':d['slides']},headers=c['owner']).status_code==200
    assert client.post(c['url']+'/restore',json={'base_revision':2,'revision':1},headers=c['owner']).status_code==200
    assert active_provider_snapshot() is None

def test_export_immutable_revision_and_zero_credits(client,case,monkeypatch,app):
    from models import Task,CreditLedger,CreditAccount
    from services.task_manager import task_manager
    from services.task_execution import task_execution_scope
    from services.provider_config import provider_snapshot_scope
    from io import BytesIO
    from pptx import Presentation
    c=case;d=initialize(client,c);pending=[]
    account=CreditAccount(user_id=c['raw']['user_id'],balance=50,reserved_balance=10,lifetime_credited=50)
    db.session.add(account);db.session.commit()
    before_account=account.to_dict()
    before=(CreditAccount.query.count(),CreditLedger.query.count())
    def capture(task_id,worker,**kwargs):
        assert not kwargs['provider_snapshot'].values
        assert kwargs['provider_snapshot'].user_id==kwargs['user_id']
        pending.append((task_id,worker,kwargs))
    monkeypatch.setattr(task_manager,'submit_task',capture)
    r=client.post(c['url']+'/export',json={'revision':1},headers=c['owner']);assert r.status_code==202,r.json
    taskid=r.json['data']['task_id']
    assert client.post(c['url']+'/export',json={'revision':1},headers=c['owner']).json['data']['task_id']==taskid
    d['slides'][0]['elements'][0]['content']=d['slides'][0]['elements'][0]['content'].replace('中文','后来')
    assert client.put(c['url'],json={'base_revision':1,'slides':d['slides']},headers=c['owner']).status_code==200
    tid,worker,kw=pending[0];snapshot=kw.pop('provider_snapshot');control=kw.pop('execution_context')
    with provider_snapshot_scope(snapshot),task_execution_scope(control):worker(tid,**kw)
    response=client.get(c['url']+'/exports/'+taskid,headers=c['owner']);assert response.status_code==200
    deck=Presentation(BytesIO(response.data));assert deck.slides[0].shapes[0].text=='中文语义编辑'
    assert client.get(c['url']+'/exports/'+taskid,headers=c['other']).status_code==404
    assert client.get(c['url']+'/export-tasks/'+taskid,headers=c['owner']).json['data']['status']=='COMPLETED'
    assert (CreditAccount.query.count(),CreditLedger.query.count())==before
    db.session.refresh(account);assert account.to_dict()==before_account

def test_cookie_csrf_and_single_user_gate(client,case,monkeypatch):
    c=case;initialize(client,c)
    login=client.post('/api/auth/login',json={'username':'editor-owner','password':'test-password'})
    assert login.status_code==200
    d=client.get(c['url']).json['data']
    assert client.put(c['url'],json={'base_revision':1,'slides':d['slides']}).status_code==403
    csrf=client.get_cookie('banana_csrf_token').value
    assert client.put(c['url'],json={'base_revision':1,'slides':d['slides']},headers={'X-CSRF-Token':csrf}).status_code==200
    monkeypatch.setenv('SINGLE_USER_MODE','true');monkeypatch.setenv('SINGLE_USER_MODE_USER_ID','not-this-user')
    assert client.get(c['url']).status_code in (403,503)

def test_asset_snapshot_freeze_and_history_access(client,case,app,tmp_path):
    from scripts.semantic_export.sample import samples
    from pathlib import Path
    c=case
    root=Path(app.config['UPLOAD_FOLDER'])/c['project_id'];root.mkdir()
    page=samples(root)[0].model_dump(mode='json');page.update(id=c['raw']['id'],project_id=c['project_id'],user_id=c['raw']['user_id'])
    for a in page['assets']:a.update(project_id=c['project_id'],user_id=c['raw']['user_id'])
    r=client.post(c['url']+'/initialize',json={'base_revision':0,'pages':[page]},headers=c['owner']);assert r.status_code==200,r.json
    d=r.json['data'];asset=d['pages'][0]['assets'][0]
    assert asset['uri'].startswith('editor-assets/')
    (root/'assets/software.png').write_bytes(b'original source overwritten')
    url=d['slides'][0]['elements'][4]['src']
    response=client.get(url,headers=c['owner']);assert response.status_code==200
    assert response.headers['X-Content-Type-Options']=='nosniff'
    assert client.get(url,headers=c['other']).status_code==404
    assert client.get(url.replace(asset['sha256'],'0'*64),headers=c['owner']).status_code==404
    assert client.put(c['url'],json={'base_revision':1,'slides':d['slides']},headers=c['owner']).status_code==200
    assert client.get(url,headers=c['owner']).status_code==200
