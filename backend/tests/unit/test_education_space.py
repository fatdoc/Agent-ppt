"""File lifecycle, ownership, retained dependencies and real export inventory."""

import io
from pathlib import Path
import pytest
from models import db, Project, ReferenceFile, Material, Task, User
from models.education_space import EducationSpaceEntry

API = "/api/competition/space"


@pytest.fixture
def edu(client, app, monkeypatch):
    monkeypatch.setitem(app.config, "APP_EDITION", "education")
    return client


def project(client, name="空间验收"):
    return client.post("/api/competition/projects", json={"name": name}).json["data"]


def upload(client, project_id=None):
    payload = {"file": (io.BytesIO("真实资料，效果待核验".encode()), "项目材料.txt")}
    if project_id:
        payload["project_id"] = project_id
    r = client.post("/api/reference-files/upload", data=payload)
    assert r.status_code == 200, r.json
    return r.json["data"]["file"]["id"]


def items(client):
    r = client.get(API)
    assert r.status_code == 200, r.json
    return r.json["data"]["items"]


def change(client, kind, ident, action, **kwargs):
    return client.patch(f"{API}/{kind}/{ident}", json={"action": action, **kwargs})


def test_space_hidden_in_general(client):
    assert client.get(API).status_code == 404


def test_classification_and_recycle_preserve_actual_file(edu, app):
    project(edu)
    ident = upload(edu)
    assert next(i for i in items(edu) if i["id"] == ident)["category"] == "temporary"
    assert change(edu, "reference", ident, "permanent").status_code == 200
    assert (
        change(edu, "reference", ident, "rename", name="保存的材料").status_code == 200
    )
    assert change(edu, "reference", ident, "trash").status_code == 200
    row = next(i for i in items(edu) if i["id"] == ident)
    assert row["category"] == "trash" and row["name"] == "保存的材料.txt"
    path = (
        Path(app.config["UPLOAD_FOLDER"])
        / db.session.get(ReferenceFile, ident).file_path
    )
    assert path.exists()
    assert edu.get(f"{API}/reference/{ident}/download").status_code == 404
    assert edu.post(f"/api/reference-files/{ident}/parse").status_code == 404
    assert ident not in [
        f["id"]
        for f in edu.get("/api/reference-files/project/all").json["data"]["files"]
    ]
    assert change(edu, "reference", ident, "restore").status_code == 200
    assert next(i for i in items(edu) if i["id"] == ident)["category"] == "permanent"
    assert (
        edu.get(f"{API}/reference/{ident}/download").data
        == "真实资料，效果待核验".encode()
    )


def test_linked_material_never_temporary_or_trash(edu):
    s = project(edu)
    ident = upload(edu, s["project_id"])
    row = next(i for i in items(edu) if i["id"] == ident)
    assert row["category"] == "permanent" and not row["can_trash"]
    assert change(edu, "reference", ident, "trash").status_code == 409
    assert change(edu, "project", s["project_id"], "trash").status_code == 200
    assert change(edu, "reference", ident, "trash").status_code == 409


def test_project_restore_and_rename_keep_canonical_content(edu):
    s = project(edu)
    pid = s["project_id"]
    assert change(edu, "project", pid, "rename", name="已重命名作品").status_code == 200
    current = edu.get(f"/api/projects/{pid}/competition").json["data"]
    assert current["revision"] == s["revision"] + 1
    assert current["content"]["profile"]["name"] == "已重命名作品"
    assert change(edu, "project", pid, "trash").status_code == 200
    assert edu.get(f"/api/projects/{pid}/competition").status_code == 404
    assert not edu.get("/api/competition/projects").json["data"]
    assert change(edu, "project", pid, "restore").status_code == 200
    assert (
        edu.get(f"/api/projects/{pid}/competition").json["data"]["content"]
        == current["content"]
    )


def test_running_project_and_foreign_resources_protected(edu):
    s = project(edu)
    db.session.add(
        Task(
            project_id=s["project_id"], task_type="GENERATE_IMAGES", status="PROCESSING"
        )
    )
    other = User(username="space-other", password_hash="unused")
    db.session.add(other)
    db.session.flush()
    reference = ReferenceFile(
        user_id=other.id,
        filename="私有.txt",
        file_path="private.txt",
        file_size=3,
        file_type="txt",
    )
    db.session.add(reference)
    db.session.commit()
    assert change(edu, "project", s["project_id"], "trash").status_code == 409
    assert change(edu, "reference", reference.id, "permanent").status_code == 404
    assert edu.get(f"{API}/reference/{reference.id}/download").status_code == 404
    assert reference.id not in [i["id"] for i in items(edu)]


def test_use_file_copies_bytes_and_binds_reference_without_moving_library(edu, app):
    s = project(edu)
    ident = upload(edu)
    change(edu, "reference", ident, "permanent")
    response = edu.post(
        f"{API}/reference/{ident}/use", json={"project_id": s["project_id"]}
    )
    assert response.status_code == 200, response.json
    clone_id = response.json["data"]["id"]
    origin, clone = db.session.get(ReferenceFile, ident), db.session.get(
        ReferenceFile, clone_id
    )
    assert origin.project_id is None and clone.project_id == s["project_id"]
    assert origin.file_path != clone.file_path
    assert (Path(app.config["UPLOAD_FOLDER"]) / origin.file_path).read_bytes() == (
        Path(app.config["UPLOAD_FOLDER"]) / clone.file_path
    ).read_bytes()
    content = edu.get(f'/api/projects/{s["project_id"]}/competition').json["data"][
        "content"
    ]
    assert (
        clone_id in content["reference_file_ids"]
        and ident not in content["reference_file_ids"]
    )
    assert change(edu, "reference", ident, "trash").status_code == 200
    assert edu.get(f"{API}/reference/{clone_id}/download").status_code == 200


def test_export_records_are_real_revision_files_and_reject_traversal(edu):
    s = project(edu)
    pid = s["project_id"]
    content = edu.get(f"/api/projects/{pid}/competition/exports/outline").data
    export = next(i for i in items(edu) if i["kind"] == "export")
    assert export["format"] == "MD" and export["status"] == "可下载"
    assert edu.get(export["download_url"]).data == content
    assert (
        edu.get(f"{API}/project/{pid}/download?filename=../../education.db").status_code
        == 404
    )
    assert change(edu, "project", pid, "trash").status_code == 200
    assert not any(i["kind"] == "export" for i in items(edu))
    assert edu.get(export["download_url"]).status_code == 404
    change(edu, "project", pid, "restore")
    assert edu.get(export["download_url"]).data == content


def test_missing_files_report_unavailable_and_general_ignores_tombstones(
    edu, app, monkeypatch
):
    s = project(edu)
    ident = upload(edu)
    path = (
        Path(app.config["UPLOAD_FOLDER"])
        / db.session.get(ReferenceFile, ident).file_path
    )
    path.unlink()
    row = next(i for i in items(edu) if i["id"] == ident)
    assert row["status"] == "文件缺失" and row["download_url"] is None
    assert edu.get(f"{API}/reference/{ident}/download").status_code == 404
    change(edu, "project", s["project_id"], "trash")
    monkeypatch.setitem(app.config, "APP_EDITION", "general")
    assert edu.get(f'/api/projects/{s["project_id"]}').status_code == 200


def test_photo_library_keeps_chinese_name_and_copies_independent_asset(edu, app):
    from PIL import Image

    s = project(edu)
    image = io.BytesIO()
    Image.new("RGB", (8, 8), "white").save(image, format="PNG")
    image.seek(0)
    response = edu.post("/api/materials/upload", data={"file": (image, "团队照片.png")})
    assert response.status_code == 201, response.json
    ident = response.json["data"]["id"]
    assert response.json["data"]["original_filename"] == "团队照片.png"
    assert change(edu, "material", ident, "permanent").status_code == 200
    assert (
        change(edu, "material", ident, "rename", name="指导老师照片").status_code == 200
    )
    use = edu.post(f"{API}/material/{ident}/use", json={"project_id": s["project_id"]})
    assert use.status_code == 200, use.json
    clone = db.session.get(Material, use.json["data"]["id"])
    original = db.session.get(Material, ident)
    assert clone.original_filename == "指导老师照片.png"
    assert original.project_id is None and clone.project_id == s["project_id"]
    assert clone.relative_path != original.relative_path
    assert change(edu, "material", ident, "trash").status_code == 200
    listed = edu.get("/api/materials?project_id=all").json["data"]["materials"]
    assert ident not in [r["id"] for r in listed] and clone.id in [
        r["id"] for r in listed
    ]
    assert edu.get(f"{API}/material/{clone.id}/download").status_code == 200


def test_malformed_operations_are_rejected_without_mutation(edu):
    s = project(edu)
    endpoint = f'{API}/project/{s["project_id"]}'
    assert edu.patch(endpoint, json=["rename"]).status_code == 400
    assert (
        edu.patch(endpoint, json={"action": "rename", "name": "bad\nname"}).status_code
        == 400
    )
    assert db.session.get(Project, s["project_id"]).project_title == "空间验收"


def test_text_preview_is_read_only_and_keeps_uploaded_content(edu):
    project(edu)
    ident = upload(edu)
    response = edu.get(f"{API}/reference/{ident}/preview")
    assert response.status_code == 200, response.json
    assert response.json["data"]["kind"] == "text"
    assert response.json["data"]["text"] == "真实资料，效果待核验"
    assert response.headers["Cache-Control"] == "private, no-store"
    assert db.session.get(ReferenceFile, ident).parse_status == "pending"
    assert Task.query.count() == 0
    change(edu, "reference", ident, "trash")
    assert edu.get(f"{API}/reference/{ident}/preview").status_code == 404


def test_image_preview_and_foreign_access(edu):
    from PIL import Image

    project(edu)
    image = io.BytesIO()
    Image.new("RGB", (32, 24), "white").save(image, format="PNG")
    image.seek(0)
    row = edu.post("/api/materials/upload", data={"file": (image, "预览.png")}).json[
        "data"
    ]
    url = f'{API}/material/{row["id"]}/preview'
    assert edu.get(url).json["data"] == {"kind": "image", "width": 32, "height": 24}
    result = edu.get(url + "?mode=image")
    assert result.status_code == 200 and result.mimetype == "image/png"
    assert result.headers["X-Content-Type-Options"] == "nosniff"
    assert Image.open(io.BytesIO(result.data)).size == (32, 24)
    other = User(username="private-preview", password_hash="unused")
    db.session.add(other)
    db.session.flush()
    db.session.get(Material, row["id"]).user_id = other.id
    db.session.commit()
    assert edu.get(url).status_code == 404
    assert edu.get(url + "?mode=image").status_code == 404


def test_pdf_preview_renders_requested_pages_without_a_model(edu):
    import fitz
    from PIL import Image

    project(edu)
    doc = fitz.open()
    doc.new_page(width=400, height=300).insert_text((25, 40), "Page one")
    doc.new_page(width=300, height=400).insert_text((25, 40), "Page two")
    raw = doc.tobytes()
    doc.close()
    ident = edu.post(
        "/api/reference-files/upload", data={"file": (io.BytesIO(raw), "预览.pdf")}
    ).json["data"]["file"]["id"]
    url = f"{API}/reference/{ident}/preview"
    assert edu.get(url).json["data"] == {"kind": "pdf", "page_count": 2}
    first = edu.get(url + "?mode=image&page=1")
    second = edu.get(url + "?mode=image&page=2")
    assert first.status_code == second.status_code == 200
    assert (
        Image.open(io.BytesIO(first.data)).size
        != Image.open(io.BytesIO(second.data)).size
    )
    assert edu.get(url + "?mode=image&page=-1").status_code == 400
    assert edu.get(url + "?mode=image&page=99").status_code in {400, 422}
    assert Task.query.count() == 0


def test_office_preview_and_corrupt_file_fallback(edu):
    from zipfile import ZipFile
    from pptx import Presentation

    project(edu)
    raw = io.BytesIO()
    with ZipFile(raw, "w") as archive:
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>成员照片与岗位说明</w:t></w:r></w:p></w:body></w:document>',
        )
    raw.seek(0)
    ident = edu.post(
        "/api/reference-files/upload", data={"file": (raw, "项目.docx")}
    ).json["data"]["file"]["id"]
    response = edu.get(f"{API}/reference/{ident}/preview").json["data"]
    assert (
        response["kind"] == "text"
        and "成员照片与岗位说明" in response["text"]
        and "不还原" in response["note"]
    )
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[1])
    slide.shapes.title.text = "技能展示"
    raw = io.BytesIO()
    ppt.save(raw)
    raw.seek(0)
    ident = edu.post(
        "/api/reference-files/upload", data={"file": (raw, "展示.pptx")}
    ).json["data"]["file"]["id"]
    assert "第 1 页" in edu.get(f"{API}/reference/{ident}/preview").json["data"]["text"]
    ident = edu.post(
        "/api/reference-files/upload",
        data={"file": (io.BytesIO(b"broken"), "损坏.pdf")},
    ).json["data"]["file"]["id"]
    r = edu.get(f"{API}/reference/{ident}/preview")
    assert r.status_code == 422 and "下载原文件" in r.json["error"]["message"]
