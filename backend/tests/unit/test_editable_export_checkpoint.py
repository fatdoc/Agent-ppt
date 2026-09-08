from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation

from services.editable_export_checkpoint import EditableExportCheckpoint
from services.export_service import ExportService
from services.image_editability import ImageEditabilityService, ServiceConfig
from services.image_editability.data_models import BBox, EditableElement, EditableImage
from services.image_editability.text_attribute_extractors import TextStyleResult, ColoredSegment


def make_image(path, color='white'):
    Image.new('RGB', (160, 90), color).save(path)
    return str(path)


def result(path, name='page'):
    box = BBox(0, 0, 100, 40)
    return EditableImage(name, path, 160, 90, elements=[
        EditableElement(name + '-text', 'text', box, box, 'hello', path)
    ], clean_background=path)


def test_checkpoint_survives_temporary_asset_removal_and_detects_corruption(tmp_path):
    source = make_image(tmp_path / 'slide.png')
    temporary = make_image(tmp_path / 'crop.png', 'red')
    cache = EditableExportCheckpoint(tmp_path / 'cache', {'depth': 1})
    key = cache.key(source)
    page = result(source)
    page.elements[0].children = [result(temporary, 'child').elements[0]]
    cache.save_analysis(key, page)
    Path(temporary).unlink()
    loaded = cache.load_analysis(key)
    assert loaded.elements[0].children[0].content == 'hello'
    assert Path(loaded.elements[0].children[0].image_path).exists()
    Path(loaded.clean_background).write_bytes(b'corrupt')
    assert cache.load_analysis(key) is None


def test_identity_invalidates_changed_image_settings_and_project(tmp_path):
    source = make_image(tmp_path / 'slide.png')
    cache = EditableExportCheckpoint(tmp_path / 'a', {'icons': True})
    key = cache.key(source)
    cache.save_analysis(key, result(source))
    assert EditableExportCheckpoint(tmp_path / 'b', {'icons': True}).load_analysis(key) is None
    assert EditableExportCheckpoint(tmp_path / 'a', {'icons': False}).key(source) != key
    make_image(Path(source), 'red')
    assert cache.key(source) != key


def test_atomic_concurrent_writes_and_style_round_trip(tmp_path):
    source = make_image(tmp_path / 'slide.png')
    cache = EditableExportCheckpoint(tmp_path / 'cache', {})
    key = cache.key(source)
    with ThreadPoolExecutor(4) as executor:
        list(executor.map(lambda i: cache.save_analysis(key, result(source, str(i))), range(8)))
    loaded = cache.load_analysis(key)
    assert loaded is not None
    style = TextStyleResult(font_color_rgb=(255, 0, 0), is_bold=True,
                            colored_segments=[ColoredSegment('hello', (0, 255, 0))])
    cache.save_styles(key, loaded.image_id, {'text': style})
    restored = cache.load_styles(key, loaded.image_id)['text']
    assert restored == style
    assert cache.load_styles(key, 'another-analysis') is None


def test_failed_export_reuses_successful_pages_and_preserves_order(tmp_path, monkeypatch):
    paths = [make_image(tmp_path / f'{i}.png') for i in range(3)]
    calls = []
    fail = {paths[1]}
    monkeypatch.setattr(ServiceConfig, 'from_defaults', lambda **kw: SimpleNamespace())
    monkeypatch.setattr(ImageEditabilityService, '__init__', lambda self, config: None)

    def analyze(self, path):
        calls.append(path)
        if path in fail:
            raise RuntimeError('simulated provider failure')
        return EditableImage(Path(path).stem, path, 160, 90, clean_background=path)

    monkeypatch.setattr(ImageEditabilityService, 'make_image_editable', analyze)
    options = dict(image_paths=paths, output_file=str(tmp_path / 'out.pptx'),
                   checkpoint_dir=str(tmp_path / 'cache'), max_workers=1)
    with pytest.raises(RuntimeError, match='simulated'):
        ExportService.create_editable_pptx_with_recursive_analysis(**options)
    assert set(calls) == set(paths)  # Includes worker success after the failure.
    calls.clear()
    fail.clear()
    messages = []
    ExportService.create_editable_pptx_with_recursive_analysis(
        **options, progress_callback=lambda step, message, percent: messages.append(message))
    assert calls == [paths[1]]
    assert any('已复用 2/3 页版面分析' in message for message in messages)
    deck = Presentation(options['output_file'])
    assert len(deck.slides) == 3
    calls.clear()
    ExportService.create_editable_pptx_with_recursive_analysis(**{
        **options, 'image_paths': [paths[2], paths[0]],
    })
    assert calls == []
    assert len(Presentation(options['output_file']).slides) == 2


def test_style_failure_retries_only_unfinished_page(tmp_path, monkeypatch):
    paths = [make_image(tmp_path / f'{i}.png') for i in range(3)]
    cache = EditableExportCheckpoint(tmp_path / 'cache', {})
    pages = [result(path, str(i)) for i, path in enumerate(paths)]
    keys = {page.image_id: cache.key(path) for page, path in zip(pages, paths)}
    calls = []
    fail = {'1'}

    def styles(images, *args, **kwargs):
        page = images[0]
        calls.append(page.image_id)
        if page.image_id in fail:
            raise RuntimeError('style failed')
        return {page.elements[0].element_id: TextStyleResult(is_bold=True)}, []

    monkeypatch.setattr(ExportService, '_batch_extract_text_styles_hybrid', styles)
    def run():
        return ExportService._extract_styles_with_checkpoints(
            pages, object(), 1, True, cache, keys, lambda *args: None)
    with pytest.raises(RuntimeError, match='style failed'):
        run()
    calls.clear()
    fail.clear()
    all_styles, errors = run()
    assert calls == ['1']
    assert len(all_styles) == 3 and not errors


def test_partial_style_results_are_not_cached(tmp_path, monkeypatch):
    path = make_image(tmp_path / 'slide.png')
    page = result(path)
    cache = EditableExportCheckpoint(tmp_path / 'cache', {})
    key = cache.key(path)
    monkeypatch.setattr(ExportService, '_batch_extract_text_styles_hybrid',
                        lambda *a, **kw: ({'page-text': TextStyleResult()}, [('page-text', 'failed')]))
    _, errors = ExportService._extract_styles_with_checkpoints(
        [page], object(), 1, False, cache, {'page': key}, lambda *a: None)
    assert errors
    assert cache.load_styles(key, 'page') is None


def test_packaging_failure_keeps_analysis_and_styles_for_retry(tmp_path, monkeypatch):
    source = make_image(tmp_path / 'slide.png')
    monkeypatch.setattr(ServiceConfig, 'from_defaults', lambda **kw: SimpleNamespace())
    monkeypatch.setattr(ImageEditabilityService, '__init__', lambda self, config: None)
    calls = {'analysis': 0, 'styles': 0}

    def analyze(self, path):
        calls['analysis'] += 1
        return result(path)

    def extract(*args, **kwargs):
        calls['styles'] += 1
        return {'page-text': TextStyleResult(is_bold=True)}, []

    monkeypatch.setattr(ImageEditabilityService, 'make_image_editable', analyze)
    monkeypatch.setattr(ExportService, '_batch_extract_text_styles_hybrid', extract)
    original = ExportService._add_editable_elements_to_slide
    def fail_build(*args, **kwargs):
        raise RuntimeError('packaging interrupted')
    monkeypatch.setattr(ExportService, '_add_editable_elements_to_slide', fail_build)
    options = dict(image_paths=[source], output_file=str(tmp_path / 'editable.pptx'),
                   checkpoint_dir=str(tmp_path / 'cache'), text_attribute_extractor=object())
    with pytest.raises(RuntimeError, match='packaging interrupted'):
        ExportService.create_editable_pptx_with_recursive_analysis(**options)
    monkeypatch.setattr(ExportService, '_add_editable_elements_to_slide', original)
    # Fully cached exports must not initialize remote providers/models at all.
    def unexpected_init(**kwargs):
        raise AssertionError('cached pages should not initialize analysis providers')
    monkeypatch.setattr(ServiceConfig, 'from_defaults', unexpected_init)
    ExportService.create_editable_pptx_with_recursive_analysis(**options)
    assert calls == {'analysis': 1, 'styles': 1}
    deck = Presentation(options['output_file'])
    assert [shape.text for shape in deck.slides[0].shapes if shape.has_text_frame] == ['hello']


def test_incomplete_analysis_is_not_reused(tmp_path, monkeypatch):
    source = make_image(tmp_path / 'slide.png')
    monkeypatch.setattr(ServiceConfig, 'from_defaults', lambda **kw: SimpleNamespace())
    monkeypatch.setattr(ImageEditabilityService, '__init__', lambda self, config: None)
    calls = []
    def analyze(self, path):
        calls.append(path)
        return EditableImage('partial', path, 160, 90, metadata={'checkpoint_complete': False})
    monkeypatch.setattr(ImageEditabilityService, 'make_image_editable', analyze)
    options = dict(image_paths=[source], checkpoint_dir=str(tmp_path / 'cache'),
                   output_file=str(tmp_path / 'out.pptx'), fail_fast=False)
    ExportService.create_editable_pptx_with_recursive_analysis(**options)
    ExportService.create_editable_pptx_with_recursive_analysis(**options)
    assert calls == [source, source]
